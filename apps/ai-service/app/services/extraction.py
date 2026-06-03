"""Real text extraction for uploaded CV files (PDF / DOCX / plain text).

Uses pdfminer.six for PDF and python-docx for DOCX. Both are pure-Python and
degrade gracefully: if a library is missing or extraction fails, we fall back
to a best-effort UTF-8 decode so the screening pipeline never hard-fails.
"""
import io
import logging

logger = logging.getLogger(__name__)


def extract_text(content: bytes, filename: str) -> str:
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            return _extract_pdf(content)
        if name.endswith(".docx"):
            return _extract_docx(content)
    except Exception as exc:  # pragma: no cover - defensive
        logger.warning("Extraction failed for %s: %s", filename, exc)

    return content.decode("utf-8", errors="ignore")


def _extract_pdf(content: bytes) -> str:
    from pdfminer.high_level import extract_text as pdf_extract_text

    text = pdf_extract_text(io.BytesIO(content)) or ""
    return text.strip()


def _extract_docx(content: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(content))
    parts = [para.text for para in document.paragraphs if para.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts).strip()
